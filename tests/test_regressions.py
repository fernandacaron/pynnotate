import builtins
import copy
import hashlib
import json
import sys
import tempfile
import types
import unittest
from collections import defaultdict
from pathlib import Path
from unittest.mock import patch

from pynnotate.main.main import (
    assign_flexible_taxon_labels,
    begin_search,
    download_sequences,
    filter_ind_by_spp,
    get_gene_extraction_failure_reason,
    make_query,
    metadata_gene_vs_sample,
    overlap_animals,
    standardize_organism,
)
from pynnotate.main import synonym_updates
from pynnotate.graphic.utils import copy_text_to_clipboard, format_elapsed_time
from scripts.build_synonym_database import build_database


class Location:
    def __init__(self, start, end, strand=1):
        self.start = start
        self.end = end
        self.strand = strand


class Feature:
    def __init__(self, gene, start, end, strand=1):
        self.type = "gene"
        self.qualifiers = {"gene": [gene]}
        self.location = Location(start, end, strand)


class Record:
    def __init__(self):
        self.id = "ABC123.1"
        self.annotations = {"original_id": "ABC123", "organism": "Example species"}
        self.features = [Feature("COI", 0, 100), Feature("COII", 90, 180)]
        self.seq = "A" * 180
        self.description = ""


class SourceFeature:
    type = "source"
    qualifiers = {
        "organism": ["Example species"],
        "geo_loc_name": ["Brazil: Rio Grande do Norte"],
    }


class RegressionTests(unittest.TestCase):
    @staticmethod
    def make_record(accession, organism, isolate, genes):
        record = Record()
        record.id = f"{accession}.1"
        record.annotations = {"organism": organism}
        source = types.SimpleNamespace(
            type="source",
            qualifiers={"organism": [organism], "isolate": [isolate]},
        )
        features = [source]
        start = 0
        for gene in genes:
            features.append(Feature(gene, start, start + 100))
            start += 100
        record.features = features
        record.seq = "A" * max(start, 100)
        record.description = ""
        return record

    def test_overlap_is_removed_and_reported(self):
        gene_dict = {
            "COI": [("Example_species_ABC123", "A" * 100, "ABC123")],
            "COII": [("Example_species_ABC123", "C" * 90, "ABC123")],
        }
        _, _, overlap_log = overlap_animals(
            Record(), gene_dict, {"COI": ["COI"], "COII": ["COII"]}
        )

        self.assertEqual(len(gene_dict["COI"][0][1]), 90)
        self.assertEqual(overlap_log[0][-1], 10)

    def test_gene_extraction_failure_reasons_are_distinct(self):
        alias_map = {"COI": ["CO1", "COX1"], "16S": ["16S rRNA"]}

        unannotated = self.make_record("NOANN", "Example species", "A", [])
        self.assertIn(
            "no annotated gene features",
            get_gene_extraction_failure_reason(unannotated, alias_map, ["COI"]),
        )

        unknown_alias = self.make_record(
            "UNKNOWN", "Example species", "B", ["cytochrome oxidase subunit one"]
        )
        self.assertIn(
            "not recognized by the selected synonym dictionary",
            get_gene_extraction_failure_reason(unknown_alias, alias_map, ["COI"]),
        )

        different_target = self.make_record(
            "OTHER", "Example species", "C", ["16S rRNA"]
        )
        self.assertIn(
            "do not match the selected target genes",
            get_gene_extraction_failure_reason(different_target, alias_map, ["COI"]),
        )

    def test_strict_mode_uses_length_then_accession_to_break_gene_count_ties(self):
        alias_map = {"COI": ["COI"], "16S": ["16S"]}

        shorter = self.make_record(
            "ZZZ200", "Example species", "shorter", ["COI", "16S"]
        )
        longer = self.make_record(
            "ZZZ100", "Example species", "longer", ["COI", "16S"]
        )
        longer.features[1].location.end = 150
        longer.features[2].location.start = 150
        longer.features[2].location.end = 300
        longer.seq = "A" * 300

        reasons = defaultdict(list)
        selected = filter_ind_by_spp(
            [shorter, longer], alias_map, {"COI", "16S"}, reasons, False, True
        )
        self.assertEqual([record.id for record in selected], ["ZZZ100.1"])
        self.assertIn("greater cumulative target-gene length", reasons["ZZZ200.1"][0])

        later_accession = self.make_record(
            "ZZZ999", "Other species", "later", ["COI", "16S"]
        )
        earlier_accession = self.make_record(
            "AAA001", "Other species", "earlier", ["COI", "16S"]
        )
        reasons = defaultdict(list)
        selected = filter_ind_by_spp(
            [later_accession, earlier_accession],
            alias_map,
            {"COI", "16S"},
            reasons,
            False,
            True,
        )
        self.assertEqual([record.id for record in selected], ["AAA001.1"])
        self.assertIn("stable GenBank-accession", reasons["ZZZ999.1"][0])

    def test_provisional_taxa_are_not_collapsed(self):
        self.assertNotEqual(
            standardize_organism("Physalaemus sp. 1"),
            standardize_organism("Physalaemus sp. 2"),
        )

    def test_cancelling_manual_download_does_not_create_output_folder(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_dir = Path(temp_dir) / "cancelled-download"
            config = {
                "email": "researcher@example.org",
                "title": "",
                "additional": "",
                "organisms": "",
                "genes": "",
                "fields": "organism",
                "output_name": temp_dir,
                "folder_name": output_dir.name,
                "include_id": True,
                "mito": False,
                "mitogenome": False,
                "ids_text": ["ABC123"],
                "chloroplast": False,
                "min_len": "",
                "max_len": "",
                "delete_unverified": False,
                "extraction": True,
                "overlap": False,
                "unique": False,
                "prioritize_more_genes": False,
                "logmissing": False,
                "org_type": "animal_mito",
                "api_key": "",
                "add_synonyms": {},
            }

            with patch.object(builtins, "input", return_value="n"):
                begin_search(config)

            self.assertFalse(output_dir.exists())

    def test_cancelling_query_download_does_not_create_output_folder(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_dir = Path(temp_dir) / "cancelled-query"
            config = {
                "email": "researcher@example.org",
                "title": "",
                "additional": "",
                "organisms": "Anura",
                "genes": "COI",
                "fields": "organism",
                "output_name": temp_dir,
                "folder_name": output_dir.name,
                "include_id": True,
                "mito": False,
                "mitogenome": False,
                "ids_text": [],
                "chloroplast": False,
                "min_len": "",
                "max_len": "",
                "delete_unverified": False,
                "extraction": True,
                "overlap": False,
                "unique": False,
                "prioritize_more_genes": False,
                "logmissing": False,
                "org_type": "animal_mito",
                "api_key": "",
                "add_synonyms": {},
            }

            with patch("pynnotate.main.main.make_query", return_value=(["ABC123"], "COI AND Anura")):
                with patch.object(builtins, "input", return_value="n"):
                    begin_search(config)

            self.assertFalse(output_dir.exists())

    def test_gui_query_confirmation_receives_complete_query(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            config = {
                "email": "researcher@example.org",
                "title": "",
                "additional": "",
                "organisms": "Anura",
                "genes": "COI",
                "fields": "organism",
                "output_name": temp_dir,
                "folder_name": "cancelled-gui-query",
                "include_id": True,
                "mito": False,
                "mitogenome": False,
                "ids_text": [],
                "chloroplast": False,
                "min_len": "",
                "max_len": "",
                "delete_unverified": False,
                "extraction": True,
                "overlap": False,
                "unique": False,
                "prioritize_more_genes": False,
                "logmissing": False,
                "org_type": "animal_mito",
                "api_key": "",
                "add_synonyms": {},
            }
            full_query = '(COI OR CO1 OR COX1) AND ("Anura"[Organism]) NOT wgs[prop] NOT tsa[prop]'
            fake_root = object()

            with patch(
                "pynnotate.main.main.make_query",
                return_value=(["ABC123"], full_query),
            ):
                with patch(
                    "pynnotate.graphic.utils.confirm_query_download",
                    return_value=False,
                ) as confirm:
                    begin_search(config, root=fake_root)

            confirm.assert_called_once_with(fake_root, 1, full_query)
            self.assertFalse((Path(temp_dir) / "cancelled-gui-query").exists())

    def test_copy_text_to_clipboard_uses_tk_clipboard(self):
        class ClipboardRoot:
            def __init__(self):
                self.value = None
                self.updated = False

            def clipboard_clear(self):
                self.value = ""

            def clipboard_append(self, value):
                self.value += value

            def update_idletasks(self):
                self.updated = True

        root = ClipboardRoot()
        copy_text_to_clipboard(root, "COI AND Anura")

        self.assertEqual(root.value, "COI AND Anura")
        self.assertTrue(root.updated)

    def test_elapsed_time_is_formatted_without_estimating_completion(self):
        self.assertEqual(format_elapsed_time(0), "00:00:00")
        self.assertEqual(format_elapsed_time(65.9), "00:01:05")
        self.assertEqual(format_elapsed_time(3661), "01:01:01")
        self.assertEqual(format_elapsed_time(-10), "00:00:00")

    def test_flexible_requires_gene_extraction(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            output_dir = Path(temp_dir) / "invalid-flexible"
            config = {
                "email": "researcher@example.org",
                "organisms": "Anura",
                "genes": "COI",
                "fields": "organism",
                "output_name": temp_dir,
                "folder_name": output_dir.name,
                "extraction": False,
                "unique": True,
            }

            with patch("pynnotate.main.main.report_message") as report:
                begin_search(config)

            self.assertFalse(output_dir.exists())
            self.assertIn("requires 'Separate gene extraction'", report.call_args.args[0])

    def test_api_key_is_forwarded_to_entrez_search(self):
        with patch("pynnotate.main.main.search_genbank", return_value=[]) as search:
            make_query(
                "researcher@example.org", {"COI": ["COI"]}, "COI", "Anura",
                False, False, False, "", "", "", "", False, "secret-key",
            )

        self.assertEqual(search.call_args.kwargs["api_key"], "secret-key")

    def test_gene_sample_matrix_and_geo_location_metadata(self):
        gene_dict = {
            "16S": [
                ("Example_species_ABC123", "A" * 50, "ABC123"),
                ("Other_species_GHI789", "C" * 50, "GHI789"),
            ],
            "COI": [("Example_species_DEF456", "G" * 50, "DEF456")],
        }
        excel_data = [
            (
                "ABC123",
                {
                    "organism": "Example species",
                    "isolate": "N88",
                    "geo_loc_name": "Brazil: Rio Grande do Norte",
                },
            ),
            (
                "DEF456",
                {
                    "organism": "Example species",
                    "isolate": "N88",
                    "geo_loc_name": "Brazil: Rio Grande do Norte",
                },
            ),
            (
                "GHI789",
                {
                    "organism": "Other species",
                    "isolate": "N88",
                    "geo_loc_name": "Brazil: Paraiba",
                },
            ),
        ]

        for prioritize in (False, True):
            with self.subTest(prioritize_more_genes=prioritize):
                with tempfile.TemporaryDirectory() as temp_dir:
                    conflicts = metadata_gene_vs_sample(
                        gene_dict,
                        excel_data,
                        ["organism", "isolate", "geo_loc_name"],
                        temp_dir,
                        {"prioritize_more_genes": prioritize},
                    )

                    matrix_path = Path(temp_dir) / "genes_matrix.xlsx"
                    self.assertTrue(matrix_path.exists())

                    from openpyxl import load_workbook

                    workbook = load_workbook(matrix_path, read_only=True)
                    try:
                        ws = workbook.active
                        headers = list(next(ws.iter_rows(values_only=True)))
                        rows = list(ws.iter_rows(min_row=2, values_only=True))
                    finally:
                        workbook.close()
                    self.assertEqual(
                        headers[:3], ["organism", "isolate", "geo_loc_name"]
                    )
                    self.assertEqual(len(rows), 2)
                    self.assertEqual(rows[0][0:3], (
                        "Example species", "N88", "Brazil: Rio Grande do Norte"
                    ))
                    self.assertEqual(rows[0][3:5], ("ABC123", "DEF456"))
                    self.assertEqual(rows[1][0:2], ("Other species", "N88"))
                    self.assertEqual(conflicts, [])

    def test_gene_sample_matrix_preserves_same_gene_conflicts(self):
        gene_dict = {
            "16S": [
                ("Example_species_ABC123", "A" * 50, "ABC123"),
                ("Example_species_DEF456", "C" * 50, "DEF456"),
            ],
        }
        excel_data = [
            ("ABC123", {"organism": "Example species", "isolate": "N88"}),
            ("DEF456", {"organism": "Example species", "isolate": "N88"}),
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            conflicts = metadata_gene_vs_sample(
                gene_dict,
                excel_data,
                ["organism", "isolate"],
                temp_dir,
                {"prioritize_more_genes": False},
            )

            from openpyxl import load_workbook

            workbook = load_workbook(
                Path(temp_dir) / "genes_matrix.xlsx", read_only=True
            )
            try:
                ws = workbook.active
                rows = list(ws.iter_rows(min_row=2, values_only=True))
            finally:
                workbook.close()
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0][2], "ABC123; DEF456")
            self.assertEqual(conflicts[0]["individual"], "N88")
            self.assertEqual(conflicts[0]["gene"], "16S")

    def test_flexible_builds_species_matrix_independently_of_heading(self):
        records = [
            self.make_record("ABC123", "Example species", "N88", ["12S"]),
            self.make_record("ABC124", "Example species", "N88", ["COI"]),
            self.make_record("ABC125", "Example species", "N90", ["16S"]),
            self.make_record("ABC126", "Example species", "N91", ["12S"]),
        ]

        fake_bio = types.ModuleType("Bio")
        fake_bio.Entrez = types.SimpleNamespace(email=None, api_key=None)
        def fake_write(records, handle, fmt):
            count = 0
            for record in records:
                handle.write(f">{record.id}\n{record.seq}\n")
                count += 1
            return count

        fake_bio.SeqIO = types.SimpleNamespace(write=fake_write)

        class FakeProgress:
            def __init__(self, *args, **kwargs):
                pass

            def update(self, value):
                pass

            def close(self):
                pass

        fake_tqdm = types.ModuleType("tqdm")
        fake_tqdm.tqdm = FakeProgress

        config = {
            "include_id": True,
            "extract_all_genes": True,
            "delete_overlap": False,
            "unique_species": True,
            "prioritize_more_genes": False,
            "min_len": 0,
            "max_len": 999999,
            "log_missing": False,
            "org_type": "animal_mito",
            "api_key": "",
            "add_synonyms": {},
            "mitogenome": False,
            "chloroplast": False,
        }
        alias_map = {"12S": ["12S"], "16S": ["16S"], "COI": ["COI"]}

        with tempfile.TemporaryDirectory() as temp_dir:
            with patch.dict(sys.modules, {"Bio": fake_bio, "tqdm": fake_tqdm}):
                with patch(
                    "pynnotate.main.main.fetch_with_retry", return_value=records
                ):
                    download_sequences(
                        "researcher@example.org",
                        [record.id for record in records],
                        temp_dir,
                        ["organism", "specimen_voucher"],
                        alias_map,
                        config,
                        "12S, 16S, COI",
                    )

            from openpyxl import load_workbook

            workbook = load_workbook(
                Path(temp_dir) / "genes_matrix.xlsx", read_only=True
            )
            try:
                ws = workbook["metadata_gene_vs_sample"]
                headers = list(next(ws.iter_rows(values_only=True)))
                rows = list(ws.iter_rows(min_row=2, values_only=True))
                provenance_ws = workbook["gene_provenance"]
                provenance_headers = list(
                    next(provenance_ws.iter_rows(values_only=True))
                )
                provenance_rows = list(
                    provenance_ws.iter_rows(min_row=2, values_only=True)
                )
            finally:
                workbook.close()

            self.assertEqual(headers, ["organism", "12S", "16S", "COI"])
            self.assertEqual(rows, [
                ("Example species", "ABC123", "ABC125", "ABC124")
            ])

            full_fasta = (Path(temp_dir) / "sequences.fasta").read_text(
                encoding="utf-8"
            )
            self.assertIn(">Example_species_ABC123", full_fasta)
            self.assertIn(">Example_species_ABC124", full_fasta)
            self.assertIn(">Example_species_ABC125", full_fasta)

            for gene in ("12S", "16S", "COI"):
                gene_fasta = (Path(temp_dir) / "genes" / f"{gene}.fasta").read_text(
                    encoding="utf-8"
                )
                self.assertIn(">Example_species", gene_fasta)
                self.assertNotIn("ABC12", gene_fasta)

            self.assertEqual(provenance_headers, [
                "organism", "gene", "GenBank_ID", "isolate", "specimen_voucher"
            ])
            self.assertEqual(len(provenance_rows), 3)
            provenance_by_accession = {row[2]: row for row in provenance_rows}
            self.assertEqual(
                provenance_by_accession["ABC123"],
                ("Example species", "12S", "ABC123", "N88", None),
            )
            self.assertEqual(
                provenance_by_accession["ABC125"],
                ("Example species", "16S", "ABC125", "N90", None),
            )

            log_text = (Path(temp_dir) / "log.txt").read_text(encoding="utf-8")
            self.assertIn("ABC126", log_text)
            self.assertIn("All genes already saved for species", log_text)
            self.assertIn("organism-only species labels", log_text)
            self.assertNotIn("[HEADER FALLBACKS]", log_text)

    def test_flexible_taxon_labels_ignore_selected_complete_fasta_heading(self):
        records = [
            self.make_record("ABC123", "Example species", "N88", ["12S"]),
            self.make_record("ABC124", "Example species", "N88", ["COI"]),
            self.make_record("ABC125", "Example species", "N90", ["16S"]),
        ]

        assign_flexible_taxon_labels(records, ["isolate"])

        self.assertEqual(
            [record.annotations["pynnotate_taxon_label"] for record in records],
            ["Example_species", "Example_species", "Example_species"],
        )

    def test_flexible_taxon_label_is_species_with_missing_organism_fallback(self):
        first = self.make_record("ABC123", "Example species", "N88", ["12S"])
        second = self.make_record("DEF456", "Other species", "N88", ["16S"])
        missing = self.make_record("GHI789", "Third species", "unused", ["COI"])
        missing.annotations.pop("organism")
        missing.features[0].qualifiers.pop("organism")

        assign_flexible_taxon_labels([first, second, missing], ["isolate"])

        self.assertEqual(
            first.annotations["pynnotate_taxon_label"], "Example_species"
        )
        self.assertEqual(
            second.annotations["pynnotate_taxon_label"], "Other_species"
        )
        self.assertEqual(missing.annotations["pynnotate_taxon_label"], "GHI789")
        self.assertEqual(
            missing.annotations["pynnotate_taxon_label_source"], "genbank_id"
        )

    def test_missing_gene_log_does_not_prevent_metadata_creation(self):
        record = Record()
        record.features = [SourceFeature(), Feature("COI", 0, 100)]
        record.seq = "A" * 100

        fake_bio = types.ModuleType("Bio")
        fake_bio.Entrez = types.SimpleNamespace(email=None, api_key=None)
        fake_bio.SeqIO = types.SimpleNamespace(write=lambda records, handle, fmt: 1)

        class FakeProgress:
            def __init__(self, *args, **kwargs):
                pass

            def update(self, value):
                pass

            def close(self):
                pass

        fake_tqdm = types.ModuleType("tqdm")
        fake_tqdm.tqdm = FakeProgress

        config = {
            "include_id": True,
            "extract_all_genes": True,
            "delete_overlap": False,
            "unique_species": False,
            "prioritize_more_genes": False,
            "min_len": 0,
            "max_len": 999999,
            "log_missing": True,
            "org_type": "animal_mito",
            "api_key": "test-key",
            "add_synonyms": {},
            "mitogenome": False,
            "chloroplast": False,
        }

        with tempfile.TemporaryDirectory() as temp_dir:
            with patch.dict(sys.modules, {"Bio": fake_bio, "tqdm": fake_tqdm}):
                with patch("pynnotate.main.main.fetch_with_retry", return_value=[record]):
                    download_sequences(
                        "researcher@example.org", ["ABC123"], temp_dir,
                        ["organism"], {"COI": ["COI"], "COII": ["COII"]},
                        config, "COI, COII",
                    )

            self.assertTrue((Path(temp_dir) / "metadata.xlsx").exists())
            from openpyxl import load_workbook

            workbook = load_workbook(
                Path(temp_dir) / "metadata.xlsx", read_only=True
            )
            try:
                ws = workbook.active
                headers = list(next(ws.iter_rows(values_only=True)))
                row = list(next(ws.iter_rows(min_row=2, values_only=True)))
            finally:
                workbook.close()
            self.assertIn("geo_loc_name", headers)
            self.assertEqual(
                row[headers.index("geo_loc_name")],
                "Brazil: Rio Grande do Norte",
            )
            log_text = (Path(temp_dir) / "log.txt").read_text(encoding="utf-8")
            self.assertIn("COII missing", log_text)
            self.assertIn("Synonym database version:", log_text)
            synonym_snapshot = json.loads(
                (Path(temp_dir) / "synonyms_used.json").read_text(encoding="utf-8")
            )
            self.assertEqual(synonym_snapshot["genome_type"], "animal_mito")
            self.assertEqual(synonym_snapshot["synonyms"]["COI"], ["COI"])


class SynonymUpdateTests(unittest.TestCase):
    def test_database_builder_includes_all_supported_genome_types(self):
        database = build_database()

        self.assertEqual(
            set(database["dictionaries"]),
            {"animal_mito", "animal_nuclear", "plant_mito", "plant_chloro"},
        )
        self.assertIn("18S", database["dictionaries"]["animal_nuclear"])

    def test_bundled_database_contains_all_supported_genome_types(self):
        database = synonym_updates.get_bundled_database()

        self.assertEqual(
            set(database["dictionaries"]),
            {"animal_mito", "animal_nuclear", "plant_mito", "plant_chloro"},
        )
        self.assertIn("COI", database["dictionaries"]["animal_mito"])
        self.assertIn("18S", database["dictionaries"]["animal_nuclear"])
        self.assertIn("ATP1", database["dictionaries"]["plant_mito"])
        self.assertIn("rbcL", database["dictionaries"]["plant_chloro"])

    def test_manual_update_installs_all_databases_and_can_be_restored(self):
        candidate = copy.deepcopy(synonym_updates.get_bundled_database())
        candidate["database_revision"] += 1
        candidate["dictionary_version"] = "test-update"
        candidate["dictionaries"]["animal_mito"]["COI"].append("TEST-ANIMAL")
        candidate["dictionaries"]["plant_mito"]["ATP1"].append("TEST-PLANT")
        candidate["dictionaries"]["plant_chloro"]["rbcL"].append("TEST-CHLORO")
        database_bytes = (
            json.dumps(candidate, ensure_ascii=False, sort_keys=True) + "\n"
        ).encode("utf-8")
        manifest = {
            "schema_version": 1,
            "database_revision": candidate["database_revision"],
            "dictionary_version": candidate["dictionary_version"],
            "database_url": "https://updates.example/synonyms_database.json",
            "sha256": hashlib.sha256(database_bytes).hexdigest(),
        }
        manifest_bytes = json.dumps(manifest).encode("utf-8")

        def fake_fetch(url, max_bytes, timeout=12):
            if url.endswith("synonyms_manifest.json"):
                return manifest_bytes
            return database_bytes

        with tempfile.TemporaryDirectory() as temp_dir:
            with patch.dict(
                "os.environ", {"PYNNOTATE_DATA_DIR": temp_dir}, clear=False
            ):
                with patch.object(synonym_updates, "_fetch_bytes", side_effect=fake_fetch):
                    update = synonym_updates.check_for_synonym_updates(
                        "https://updates.example/synonyms_manifest.json"
                    )
                self.assertTrue(update["update_available"])
                info = synonym_updates.install_synonym_update(update)
                self.assertEqual(info["source"], "downloaded")
                self.assertIn(
                    "TEST-ANIMAL",
                    synonym_updates.get_official_alias_map("animal_mito")["COI"],
                )
                self.assertIn(
                    "TEST-PLANT",
                    synonym_updates.get_official_alias_map("plant_mito")["ATP1"],
                )
                self.assertIn(
                    "TEST-CHLORO",
                    synonym_updates.get_official_alias_map("plant_chloro")["rbcL"],
                )
                self.assertTrue(synonym_updates.restore_bundled_synonym_database())
                self.assertEqual(
                    synonym_updates.get_synonym_database_info()["source"], "bundled"
                )

    def test_checksum_mismatch_does_not_install_update(self):
        manifest = {
            "schema_version": 1,
            "database_revision": 999,
            "dictionary_version": "invalid",
            "database_url": "https://updates.example/synonyms_database.json",
            "sha256": "0" * 64,
        }

        def fake_fetch(url, max_bytes, timeout=12):
            if url.endswith("synonyms_manifest.json"):
                return json.dumps(manifest).encode("utf-8")
            return b"{}"

        with tempfile.TemporaryDirectory() as temp_dir:
            with patch.dict(
                "os.environ", {"PYNNOTATE_DATA_DIR": temp_dir}, clear=False
            ):
                with patch.object(synonym_updates, "_fetch_bytes", side_effect=fake_fetch):
                    with self.assertRaises(synonym_updates.SynonymDatabaseError):
                        synonym_updates.check_for_synonym_updates(
                            "https://updates.example/synonyms_manifest.json"
                        )
                self.assertFalse(synonym_updates.get_cached_database_path().exists())


if __name__ == "__main__":
    unittest.main()
